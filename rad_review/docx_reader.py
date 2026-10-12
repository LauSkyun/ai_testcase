from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable
from zipfile import ZipFile

try:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError as exc:  # pragma: no cover - exercised by real CLI envs
    raise RuntimeError("Missing dependency: python-docx. Install with `pip install -r requirements.txt`.") from exc


DOCUMENT_KINDS = ("requirement", "numeric", "mixed", "unknown")
SORT_PRIORITY = {"requirement": 0, "numeric": 1, "mixed": 2, "unknown": 3}
GAME_NAME_STOP_WORDS = (
    "需求文档",
    "数值文档",
    "需求",
    "数值",
    "文档",
    "游戏",
    "玩法",
    "配置",
    "prd",
    "musee",
)
REQUIREMENT_KEYWORDS = (
    "需求",
    "prd",
    "玩法",
    "功能",
    "流程",
    "界面",
    "交互",
)
NUMERIC_KEYWORDS = (
    "数值",
    "配置",
    "倍率",
    "概率",
    "rtp",
    "奖金",
    "公式",
    "赔率",
)
MIXED_REQUIREMENT_HINTS = (
    "系统文档",
    "界面布局",
    "游戏流程",
    "功能",
)
MIXED_NUMERIC_HINTS = (
    "奖金计算",
    "倍率设定",
    "抽奖规则",
    "数值",
    "配置",
)


@dataclass
class ParagraphBlock:
    index: int
    text: str
    heading_path: list[str] = field(default_factory=list)


@dataclass
class TableBlock:
    index: int
    rows: list[list[str]]
    markdown: str
    heading_path: list[str] = field(default_factory=list)


@dataclass
class ImageBlock:
    index: int
    source_name: str
    extracted_path: str
    status: str = "extracted"
    analysis: str = ""
    error: str = ""


@dataclass
class ParsedDocument:
    path: str
    name: str
    kind: str
    paragraphs: list[ParagraphBlock]
    tables: list[TableBlock]
    images: list[ImageBlock]
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    def compact_summary(self) -> dict:
        return {
            "name": self.name,
            "kind": self.kind,
            "paragraph_count": len(self.paragraphs),
            "table_count": len(self.tables),
            "image_count": len(self.images),
            "errors": self.errors,
        }


def classify_document(path: Path, paragraphs: list[ParagraphBlock] | None = None) -> str:
    name = path.stem.lower()
    name_requirement = any(keyword in name for keyword in REQUIREMENT_KEYWORDS)
    name_numeric = any(keyword in name for keyword in NUMERIC_KEYWORDS)
    if name_requirement and name_numeric:
        return "mixed"
    if name_numeric:
        return "numeric"
    if name_requirement:
        return "requirement"

    sample_parts: list[str] = []
    for paragraph in paragraphs or []:
        sample_parts.append(paragraph.text)
        if len(sample_parts) >= 40:
            break
    sample_text = "\n".join(sample_parts).lower()
    requirement_hits = keyword_hits(sample_text, REQUIREMENT_KEYWORDS)
    numeric_hits = keyword_hits(sample_text, NUMERIC_KEYWORDS)

    if requirement_hits and numeric_hits:
        if any(hint.lower() in sample_text for hint in MIXED_REQUIREMENT_HINTS) and any(
            hint.lower() in sample_text for hint in MIXED_NUMERIC_HINTS
        ):
            return "mixed"
        if requirement_hits >= 2 and numeric_hits >= 2:
            return "mixed"
    if numeric_hits >= 2 and requirement_hits == 0:
        return "numeric"
    if requirement_hits >= 2 and numeric_hits == 0:
        return "requirement"
    if requirement_hits and numeric_hits:
        return "mixed"
    return "unknown"


def discover_docx(input_path: Path) -> list[Path]:
    if input_path.is_file():
        if input_path.suffix.lower() != ".docx":
            raise ValueError(f"Input file is not .docx: {input_path}")
        return [input_path]
    if not input_path.exists():
        raise FileNotFoundError(f"Input path does not exist: {input_path}")
    docs = sorted(p for p in input_path.rglob("*.docx") if not p.name.startswith("~$"))
    if not docs:
        raise FileNotFoundError(f"No .docx files found under: {input_path}")
    return sorted(docs, key=document_sort_key)


def discover_docx_for_game(input_path: Path, game_name: str) -> list[Path]:
    docs = discover_docx(input_path)
    query = normalize_game_query(game_name)
    if not query:
        raise ValueError("--game must not be empty.")
    matched = [path for path in docs if doc_matches_game(path, query, input_path)]
    if not matched:
        available = "、".join(sorted(group_docx_by_game(docs, input_path))) or "无"
        raise FileNotFoundError(f"No .docx files matched game '{game_name}'. Available groups: {available}")
    return sorted(matched, key=document_sort_key)


def group_docx_by_game(paths: list[Path], root_dir: Path | None = None) -> dict[str, list[Path]]:
    groups: dict[str, list[Path]] = {}
    for path in paths:
        key = infer_game_key(path, root_dir)
        groups.setdefault(key, []).append(path)
    return {
        key: sorted(items, key=document_sort_key)
        for key, items in sorted(groups.items(), key=lambda item: item[0].lower())
    }


def document_sort_key(path: Path) -> tuple[int, str]:
    kind = classify_document(path)
    priority = SORT_PRIORITY[kind]
    if kind == "unknown" and not any(keyword in path.stem.lower() for keyword in NUMERIC_KEYWORDS):
        priority = SORT_PRIORITY["requirement"]
    return priority, path.name


def infer_game_key(path: Path, root_dir: Path | None = None) -> str:
    folder_key = infer_game_key_from_parent_folder(path, root_dir)
    if folder_key:
        return folder_key

    stem = path.stem.strip()
    lower = stem.lower()
    ascii_tokens = [token for token in re.findall(r"[a-z0-9]+", lower) if token not in GAME_NAME_STOP_WORDS]
    if ascii_tokens:
        return ascii_tokens[-1]

    cleaned = stem
    for word in GAME_NAME_STOP_WORDS:
        cleaned = re.sub(re.escape(word), "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"[\s_\-（）()【】\[\]·.]+", "", cleaned).strip()
    return cleaned or stem


def normalize_game_query(value: str) -> str:
    return re.sub(r"[\s_\-（）()【】\[\]·.]+", "", value.strip().lower())


def doc_matches_game(path: Path, normalized_query: str, root_dir: Path | None = None) -> bool:
    stem = normalize_game_query(path.stem)
    key = normalize_game_query(infer_game_key(path, root_dir))
    folder = normalize_game_query(infer_game_key_from_parent_folder(path, root_dir))
    return (
        normalized_query in stem
        or normalized_query in key
        or key in normalized_query
        or (folder and (normalized_query in folder or folder in normalized_query))
    )


def infer_game_key_from_parent_folder(path: Path, root_dir: Path | None = None) -> str:
    if root_dir is None:
        return ""
    try:
        relative = path.resolve().relative_to(root_dir.resolve())
    except ValueError:
        return ""
    if len(relative.parts) < 2:
        return ""
    folder_name = relative.parts[0].strip()
    return folder_name or ""


def parse_docx(path: Path, image_output_dir: Path) -> ParsedDocument:
    image_output_dir.mkdir(parents=True, exist_ok=True)
    document = Document(str(path))
    paragraphs: list[ParagraphBlock] = []
    tables: list[TableBlock] = []
    errors: list[str] = []
    heading_stack: list[str] = []
    paragraph_index = 0
    table_index = 0

    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            text = normalize_text(block.text)
            if not text:
                continue
            heading_level = detect_heading_level(text)
            if heading_level:
                heading_stack = heading_stack[: heading_level - 1]
                heading_stack.append(text)
            paragraph_index += 1
            paragraphs.append(ParagraphBlock(paragraph_index, text, list(heading_stack)))
        elif isinstance(block, Table):
            table_index += 1
            rows = extract_table_rows(block)
            tables.append(TableBlock(table_index, rows, table_to_markdown(rows), list(heading_stack)))

    images = extract_images(path, image_output_dir)
    return ParsedDocument(
        path=str(path),
        name=path.name,
        kind=classify_document(path, paragraphs),
        paragraphs=paragraphs,
        tables=tables,
        images=images,
        errors=errors,
    )


def iter_block_items(document) -> Iterable[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def normalize_text(text: str) -> str:
    return re.sub(r"[ \t\r\f\v]+", " ", text.replace("\u3000", " ")).strip()


def detect_heading_level(text: str) -> int | None:
    stripped = text.strip()
    match = re.match(r"^(\d+(?:[.、]\d+)*)(?:[.、)\s]|$)", stripped)
    if match:
        token = match.group(1).replace("、", ".")
        return max(1, min(6, token.count(".") + 1))
    if re.match(r"^第[一二三四五六七八九十百千]+[章节条]", stripped):
        return 1
    return None


def extract_table_rows(table: Table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([normalize_text(cell.text).replace("\n", " / ") for cell in row.cells])
    return rows


def table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]

    def cell(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", "<br>")

    header = "| " + " | ".join(cell(v) for v in normalized[0]) + " |"
    sep = "| " + " | ".join("---" for _ in range(width)) + " |"
    body = ["| " + " | ".join(cell(v) for v in row) + " |" for row in normalized[1:]]
    return "\n".join([header, sep, *body])


def extract_images(path: Path, image_output_dir: Path) -> list[ImageBlock]:
    images: list[ImageBlock] = []
    safe_doc_name = sanitize_filename(path.stem)
    doc_image_dir = image_output_dir / safe_doc_name
    doc_image_dir.mkdir(parents=True, exist_ok=True)
    with ZipFile(path) as archive:
        media_names = [name for name in archive.namelist() if name.startswith("word/media/")]
        for index, media_name in enumerate(media_names, start=1):
            suffix = Path(media_name).suffix or ".bin"
            output_path = doc_image_dir / f"image{index}{suffix.lower()}"
            output_path.write_bytes(archive.read(media_name))
            images.append(ImageBlock(index=index, source_name=media_name, extracted_path=str(output_path)))
    return images


def sanitize_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip()
    return value or "document"


def keyword_hits(text: str, keywords: tuple[str, ...]) -> int:
    return sum(1 for keyword in keywords if keyword in text)


def write_manifest(documents: list[ParsedDocument], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "documents": [doc.to_dict() for doc in documents],
        "summary": build_collection_summary(documents),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def build_collection_summary(documents: list[ParsedDocument]) -> dict:
    kinds = {kind: sum(1 for doc in documents if doc.kind == kind) for kind in DOCUMENT_KINDS}
    has_requirement = kinds["requirement"] > 0 or kinds["mixed"] > 0
    has_numeric = kinds["numeric"] > 0 or kinds["mixed"] > 0
    return {
        "document_count": len(documents),
        "kinds": kinds,
        "has_requirement_doc": has_requirement,
        "has_numeric_doc": has_numeric,
        "missing_companion_docs": missing_companion_docs(has_requirement, has_numeric),
        "total_paragraphs": sum(len(doc.paragraphs) for doc in documents),
        "total_tables": sum(len(doc.tables) for doc in documents),
        "total_images": sum(len(doc.images) for doc in documents),
    }


def missing_companion_docs(has_requirement_doc: bool, has_numeric_doc: bool) -> list[str]:
    missing: list[str] = []
    if not has_requirement_doc:
        missing.append("需求文档")
    if not has_numeric_doc:
        missing.append("数值文档")
    return missing
